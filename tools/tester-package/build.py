"""build.py: builds the Nyarlathotep tester package from rows.py (standard: _shared/TESTING-STANDARDS.md).

usage:
  python tools/tester-package/build.py            build xlsx, pdf, csv and html into the gitignored output folder,
                                                  then verify them (every row, title and step in every artefact;
                                                  sheet XML in schema order; merged prose rows have heights;
                                                  dropdowns valid; formulas are formulas)
  python tools/tester-package/build.py --prove    break each invariant on purpose and confirm the build refuses it

Output: REFERENCE DOCUMENTATION/TESTING DOCUMENTATION/<product>-<round>.{xlsx,pdf,csv,html} (gitignored: returned
workbooks carry tester details). Send the tester the .xlsx and the .pdf only.
"""
import copy, csv, html, math, os, re, sys, zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(REPO, "REFERENCE DOCUMENTATION", "TESTING DOCUMENTATION")
sys.path.insert(0, HERE)
import rows as R  # noqa: E402


class BuildError(Exception):
    pass


ALLOWED = {"id", "status", "area", "title", "summary", "howToTest", "needs", "mostWanted", "deliberate",
           "knownLimit", "weCouldNotVerify"}
REQUIRED = {"id", "status", "area", "title", "summary", "howToTest", "mostWanted"}
CAT_ORDER = [c for c, _ in R.CATEGORIES]
CAT_LABEL = dict(R.CATEGORIES)

TAB = {"Start here": "1F2937", "What changed": "6D28D9", "Known limits": "6B7280", "Test plan": "1D4ED8",
       "Test steps": "0F766E", "Anything new": "B91C1C", "Progress": "15803D"}
RESULT_FILL = {"Pass": "C6F6D5", "Partial pass": "FEF3C7", "Fail": "FECACA", "Blocked": "E5E7EB",
               "Not tested yet": "FFFFFF", "OK": "C6F6D5", "Problem": "FECACA", "Skipped": "E5E7EB"}
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
INPUT_HEAD_FILL = PatternFill("solid", fgColor="92400E")
BAND_FILL = PatternFill("solid", fgColor="E0E7FF")
INPUT_FILL = PatternFill("solid", fgColor="FFFBEB")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINE_PT = 15


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def lines_needed(text, width_chars):
    n = 0
    for para in str(text).split("\n"):
        n += max(1, math.ceil(len(para) / max(1, width_chars)))
    return n


# ---- data checks ------------------------------------------------------------------------------------------------

def checked_rows(rows):
    seen = set()
    for r in rows:
        extra = set(r) - ALLOWED
        if extra:
            raise BuildError(f"row {r.get('id')}: key(s) {sorted(extra)} are not rendered by build.py (refused, §6)")
        missing = REQUIRED - set(r)
        if missing:
            raise BuildError(f"row {r.get('id')}: missing required {sorted(missing)}")
        if r["status"] not in CAT_ORDER:
            raise BuildError(f"row {r['id']}: status {r['status']!r} is not a category")
        if r["status"] == "explore" and not r.get("weCouldNotVerify"):
            raise BuildError(f"row {r['id']}: an explore row needs weCouldNotVerify")
        if r["id"] in seen:
            raise BuildError(f"row id {r['id']} repeats")
        seen.add(r["id"])
        if not r["howToTest"]:
            raise BuildError(f"row {r['id']}: no steps")
    return sorted(rows, key=lambda r: CAT_ORDER.index(r["status"]))


def dropdown(options, sqref, ws):
    for o in options:
        if "," in o:
            raise BuildError(f"{ws.title}: dropdown option {o!r} contains a comma (Excel splits it)")
    formula = '"' + ",".join(options) + '"'
    if len(formula) > 255:
        raise BuildError(f"{ws.title}: dropdown list is {len(formula)} characters (max 255)")
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.add(sqref)
    ws.add_data_validation(dv)


def result_colours(ws, rng, first_cell, values):
    for v in values:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first_cell}="{v}"'],
                                                       fill=PatternFill("solid", fgColor=RESULT_FILL[v])))


def header(ws, row, names, input_from=None, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = WHITE_BOLD
        c.fill = INPUT_HEAD_FILL if input_from and i >= input_from else HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def prose(ws, row, text, cols=6, width_chars=None, bold=False, size=11, fill=None):
    """A merged prose row across `cols` columns with a computed height (§8: merged wrap rows clip otherwise)."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row=row, column=1, value=text)
    c.alignment = WRAP
    c.font = Font(bold=bold, size=size)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if width_chars is None:
        width_chars = int(sum((ws.column_dimensions[get_column_letter(i)].width or 10) for i in range(1, cols + 1)) * 1.05)
    ws.row_dimensions[row].height = lines_needed(text, width_chars) * (LINE_PT + (3 if size > 11 else 0)) + 4
    return row + 1


# ---- workbook ---------------------------------------------------------------------------------------------------

def build_workbook(rows, path):
    wb = Workbook()
    n_admin = sum(1 for r in rows if "admin" in (r.get("needs") or "").lower())
    access = R.ACCESS.format(admin_rows=n_admin, all_rows=len(rows))

    # 1. Start here
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = TAB[ws.title]
    for col, w in zip("ABCDEF", (34, 30, 30, 22, 22, 22)):
        ws.column_dimensions[col].width = w
    r = prose(ws, 1, R.ROUND["title"], bold=True, size=14)
    r = prose(ws, r, f"Build: {R.ROUND['build']}. Prepared {R.ROUND['date']}.")
    r += 1
    r = prose(ws, r, "Tester details — one column each. The line under each field says why we ask.", bold=True)
    hdr = ["Field"] + R.TESTERS + ["Why we ask"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = INPUT_HEAD_FILL if 1 < i <= 1 + len(R.TESTERS) else HEAD_FILL
    r += 1
    why_col = 2 + len(R.TESTERS)
    for field, why in R.COVER_FIELDS:
        ws.cell(row=r, column=1, value=field).alignment = WRAP
        for t in range(len(R.TESTERS)):
            c = ws.cell(row=r, column=2 + t)
            c.fill = INPUT_FILL
            c.border = BOX
        ws.cell(row=r, column=why_col, value=why).alignment = WRAP
        ws.row_dimensions[r].height = lines_needed(why, 22) * LINE_PT + 4
        r += 1
    r += 1
    r = prose(ws, r, "Where your last round's findings went", bold=True)
    r = prose(ws, r, R.LAST_ROUND)
    r = prose(ws, r, "What we could not check ourselves, and why", bold=True)
    for t in R.COULD_NOT_CHECK:
        r = prose(ws, r, "• " + t)
    r = prose(ws, r, "Your access", bold=True)
    r = prose(ws, r, access)
    r += 1
    r = prose(ws, r, "How this workbook works", bold=True)
    for name, what in [("What changed", "read first: every row, what it is and how to check it. Nothing to type."),
                       ("Known limits", "what will look like a fault and is not. Read before starting."),
                       ("Test plan", "one line per test: your result goes here."),
                       ("Test steps", "every step of every test, each with OK, Problem or Skipped."),
                       ("Anything new", "anything you notice that is not on a row."),
                       ("Progress", "counts itself. Nothing to fill in.")]:
        r = prose(ws, r, f"• {name}: {what}")
    r = prose(ws, r, "Results (Test plan)", bold=True)
    for v, what in R.RESULTS:
        r = prose(ws, r, f"• {v}: {what}", fill=RESULT_FILL[v])
    r = prose(ws, r, "Step results (Test steps)", bold=True)
    for v, what in R.STEP_RESULTS:
        r = prose(ws, r, f"• {v}: {what}", fill=RESULT_FILL[v])
    r = prose(ws, r, "Row categories, in the order you meet them", bold=True)
    for c, label in R.CATEGORIES:
        if any(x["status"] == c for x in rows):
            r = prose(ws, r, f"• {label}")
    r = prose(ws, r, "Reporting rules", bold=True)
    for t in R.RULES:
        r = prose(ws, r, "• " + t)

    # 2. What changed
    ws = wb.create_sheet("What changed")
    ws.sheet_properties.tabColor = TAB[ws.title]
    header(ws, 1, ["Row id", "Area", "What it is", "What was missing before", "What it does now", "How to check it"],
           widths=[9, 18, 40, 34, 40, 40])
    ws.freeze_panes = "B2"
    r = 2
    for cat in CAT_ORDER:
        group = [x for x in rows if x["status"] == cat]
        if not group:
            continue
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value=CAT_LABEL[cat])
        c.font = Font(bold=True)
        c.fill = BAND_FILL
        ws.row_dimensions[r].height = LINE_PT + 6
        r += 1
        for x in group:
            missing = "New — this did not exist on the server before." if cat == "new" else \
                ("Not new; never looked at by a person." if cat == "explore" else "A measurement, not a change.")
            vals = [x["id"], x["area"], x["title"], missing, x["summary"], x["howToTest"][0].split("->")[0].strip()]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.alignment = WRAP
                cell.border = BOX
            ws.row_dimensions[r].height = max(lines_needed(vals[2], 40), lines_needed(vals[4], 40),
                                              lines_needed(vals[5], 40)) * LINE_PT + 4
            r += 1

    # 3. Known limits
    ws = wb.create_sheet("Known limits")
    ws.sheet_properties.tabColor = TAB[ws.title]
    header(ws, 1, ["What we did not change", "Why"], widths=[60, 70])
    limits = list(R.KNOWN_LIMITS)
    for x in rows:
        for k in ("deliberate", "knownLimit"):
            if x.get(k):
                limits.append((f"{x['id']} ({x['title']}): {x[k]}", "Deliberate." if k == "deliberate" else
                               "Known; not this round's work."))
    for i, (what, why) in enumerate(limits, 2):
        for j, v in enumerate((what, why), 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = WRAP
            c.border = BOX
        ws.row_dimensions[i].height = max(lines_needed(what, 58), lines_needed(why, 68)) * LINE_PT + 4

    # 4. Test plan
    ws = wb.create_sheet("Test plan")
    ws.sheet_properties.tabColor = TAB[ws.title]
    read_cols = ["Row id", "Kind of row", "Area", "What to test", "You will need", "Steps", "We most want to know"]
    input_cols = ["Assigned to", "Result", "If not a clean pass — what is still wrong", "Tested at",
                  "PC + connection", "What you saw", "What you expected", "How it felt", "Evidence (screenshot name or chat line)"]
    header(ws, 1, read_cols + input_cols, input_from=len(read_cols) + 1,
           widths=[8, 22, 16, 42, 34, 7, 34, 16, 14, 34, 16, 18, 34, 34, 20, 28])
    ws.freeze_panes = "B2"
    first = 2
    for i, x in enumerate(rows, first):
        vals = [x["id"], CAT_LABEL[x["status"]], x["area"], x["title"], x.get("needs", ""), len(x["howToTest"]),
                x["mostWanted"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = WRAP
            c.border = BOX
        for j in range(len(read_cols) + 1, len(read_cols) + len(input_cols) + 1):
            c = ws.cell(row=i, column=j)
            c.fill = INPUT_FILL
            c.border = BOX
            c.alignment = WRAP
        ws.cell(row=i, column=len(read_cols) + 1, value="Not assigned yet")
        ws.cell(row=i, column=len(read_cols) + 2, value="Not tested yet")
        ws.row_dimensions[i].height = max(lines_needed(vals[3], 40), lines_needed(vals[4], 32),
                                          lines_needed(vals[6], 32)) * LINE_PT + 4
    last = first + len(rows) - 1
    ac = get_column_letter(len(read_cols) + 1)
    rc = get_column_letter(len(read_cols) + 2)
    dropdown(R.TESTERS + ["Both — coordinated", "Not assigned yet"], f"{ac}{first}:{ac}{last}", ws)
    dropdown([v for v, _ in R.RESULTS], f"{rc}{first}:{rc}{last}", ws)
    result_colours(ws, f"{rc}{first}:{rc}{last}", f"{rc}{first}", [v for v, _ in R.RESULTS])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(read_cols) + len(input_cols))}{last}"

    # 5. Test steps
    ws = wb.create_sheet("Test steps")
    ws.sheet_properties.tabColor = TAB[ws.title]
    header(ws, 1, ["Row id", "Step", "Do this", "You should see", "Result", "Notes"], input_from=5,
           widths=[8, 6, 60, 50, 12, 40])
    ws.freeze_panes = "B2"
    r = 2
    step_cells = []
    for x in rows:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value=f"{x['id']} — {x['title']}")
        c.font = Font(bold=True)
        c.fill = BAND_FILL
        c.alignment = WRAP
        ws.row_dimensions[r].height = lines_needed(c.value, 170) * LINE_PT + 6
        r += 1
        for k, step in enumerate(x["howToTest"], 1):
            do, _, see = step.partition("->")
            vals = [x["id"], k, do.strip(), see.strip()]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=j, value=v)
                cell.alignment = WRAP
                cell.border = BOX
            for j in (5, 6):
                cell = ws.cell(row=r, column=j)
                cell.fill = INPUT_FILL
                cell.border = BOX
            step_cells.append(r)
            ws.row_dimensions[r].height = max(lines_needed(vals[2], 58), lines_needed(vals[3], 48)) * LINE_PT + 4
            r += 1
    dropdown([v for v, _ in R.STEP_RESULTS], f"E2:E{r - 1}", ws)
    result_colours(ws, f"E2:E{r - 1}", "E2", [v for v, _ in R.STEP_RESULTS])
    ws.auto_filter.ref = f"A1:F{r - 1}"

    # 6. Anything new
    ws = wb.create_sheet("Anything new")
    ws.sheet_properties.tabColor = TAB[ws.title]
    for col, w in zip("ABCDEFGH", (40, 20, 40, 30, 30, 14, 20, 12)):
        ws.column_dimensions[col].width = w
    prose(ws, 1, "Anything you notice that is not on the other sheets. Small things are welcome — a line in chat that "
                 "reads wrong, a monster in an odd place, a moment of lag. If you are not sure it is a fault, put it "
                 "here anyway.", cols=8)
    header(ws, 2, ["What happened", "Where (region or place)", "Steps you took", "What you expected",
                   "What you got instead", "How bad", "PC / connection", "Date"])
    ws.column_dimensions["A"].width = 40
    for i in range(3, 25):
        for j in range(1, 9):
            c = ws.cell(row=i, column=j)
            c.fill = INPUT_FILL
            c.border = BOX
            c.alignment = WRAP
    dropdown(["Annoying", "Gets in the way", "Breaks the event", "Breaks the server"], "F3:F24", ws)
    ws.freeze_panes = "A3"

    # 7. Progress
    ws = wb.create_sheet("Progress")
    ws.sheet_properties.tabColor = TAB[ws.title]
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 12
    plan_res = f"'Test plan'!${rc}${first}:${rc}${last}"
    plan_ass = f"'Test plan'!${ac}${first}:${ac}${last}"
    plan_cat = f"'Test plan'!$B${first}:$B${last}"
    r = 1
    ws.cell(row=r, column=1, value="By result (counts itself from Test plan; nothing to fill in)").font = Font(bold=True)
    r += 1
    for v, _ in R.RESULTS:
        ws.cell(row=r, column=1, value=v).fill = PatternFill("solid", fgColor=RESULT_FILL[v])
        ws.cell(row=r, column=2, value=f'=COUNTIF({plan_res},"{v}")')
        r += 1
    ws.cell(row=r, column=1, value="Rows in the plan")
    ws.cell(row=r, column=2, value=f"=COUNTA('Test plan'!$A${first}:$A${last})")
    r += 2
    ws.cell(row=r, column=1, value="By assignee").font = Font(bold=True)
    r += 1
    for t in R.TESTERS + ["Both — coordinated", "Not assigned yet"]:
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=f'=COUNTIF({plan_ass},"{t}")')
        r += 1
    ws.cell(row=r, column=1, value="Left blank")
    ws.cell(row=r, column=2, value=f"=COUNTBLANK({plan_ass})")
    r += 2
    ws.cell(row=r, column=1, value="By kind of row").font = Font(bold=True)
    r += 1
    for c, label in R.CATEGORIES:
        if any(x["status"] == c for x in rows):
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=f'=COUNTIF({plan_cat},"{label}")')
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Steps marked (Test steps)").font = Font(bold=True)
    r += 1
    for v, _ in R.STEP_RESULTS:
        ws.cell(row=r, column=1, value=v).fill = PatternFill("solid", fgColor=RESULT_FILL[v])
        ws.cell(row=r, column=2, value=f"=COUNTIF('Test steps'!$E$2:$E${max(step_cells)},\"{v}\")")
        r += 1

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


# ---- the other artefacts ----------------------------------------------------------------------------------------

def build_csv(rows, path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "status", "area", "title", "summary", "needs", "mostWanted", "deliberate", "knownLimit",
                    "weCouldNotVerify", "steps"])
        for x in rows:
            w.writerow([x["id"], x["status"], x["area"], x["title"], x["summary"], x.get("needs", ""), x["mostWanted"],
                        x.get("deliberate", ""), x.get("knownLimit", ""), x.get("weCouldNotVerify", ""),
                        " | ".join(x["howToTest"])])


def build_html(rows, path):
    e = html.escape
    parts = [f"<!doctype html><meta charset=utf-8><title>{e(R.PRODUCT)} tester package</title>",
             "<style>body{font-family:system-ui;max-width:900px;margin:2rem auto;padding:0 1rem;line-height:1.45}"
             "h2{border-bottom:2px solid #1d4ed8}.row{border:1px solid #ddd;padding:.6rem 1rem;margin:.8rem 0}"
             ".meta{color:#555;font-size:.9em}</style>",
             f"<h1>{e(R.ROUND['title'])}</h1><p>{e(R.ROUND['build'])}</p>"]
    for cat in CAT_ORDER:
        group = [x for x in rows if x["status"] == cat]
        if not group:
            continue
        parts.append(f"<h2>{e(CAT_LABEL[cat])}</h2>")
        for x in group:
            parts.append(f"<div class=row><h3>{e(x['id'])} · {e(x['title'])}</h3><p class=meta>{e(x['area'])}"
                         f"{' · You will need: ' + e(x['needs']) if x.get('needs') else ''}</p><p>{e(x['summary'])}</p><ol>"
                         + "".join(f"<li>{e(s)}</li>" for s in x["howToTest"]) +
                         f"</ol><p><b>We most want to know:</b> {e(x['mostWanted'])}</p>"
                         + "".join(f"<p><i>{k}:</i> {e(x[k])}</p>" for k in ("deliberate", "knownLimit",
                                                                            "weCouldNotVerify") if x.get(k))
                         + "</div>")
    parts.append("<h2>Known limits</h2><ul>" + "".join(f"<li><b>{e(a)}</b> {e(b)}</li>" for a, b in R.KNOWN_LIMITS)
                 + "</ul>")
    open(path, "w", encoding="utf-8").write("".join(parts))


def build_pdf(rows, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, PageBreak
    from xml.sax.saxutils import escape as x_
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=13.5)
    small = ParagraphStyle("s", parent=body, fontSize=9, textColor="#444444")
    h1, h2, h3 = ss["Title"], ss["Heading2"], ss["Heading3"]
    n_admin = sum(1 for r in rows if "admin" in (r.get("needs") or "").lower())
    story = [Paragraph(x_(R.ROUND["title"]), h1), Paragraph(x_(f"Build: {R.ROUND['build']}. Prepared {R.ROUND['date']}."), body),
             Paragraph("This is the reading copy. Fill in your results in the workbook (.xlsx).", body), Spacer(1, 6),
             Paragraph("Where your last round's findings went", h2), Paragraph(x_(R.LAST_ROUND), body),
             Paragraph("What we could not check ourselves, and why", h2),
             ListFlowable([ListItem(Paragraph(x_(t), body)) for t in R.COULD_NOT_CHECK], bulletType="bullet"),
             Paragraph("Your access", h2),
             Paragraph(x_(R.ACCESS.format(admin_rows=n_admin, all_rows=len(rows))), body),
             Paragraph("Results", h2),
             ListFlowable([ListItem(Paragraph(x_(f"{v}: {w}"), body)) for v, w in R.RESULTS + R.STEP_RESULTS],
                          bulletType="bullet"),
             Paragraph("Reporting rules", h2),
             ListFlowable([ListItem(Paragraph(x_(t), body)) for t in R.RULES], bulletType="bullet"),
             Paragraph("Known limits", h2)]
    story.append(ListFlowable([ListItem(Paragraph(f"<b>{x_(a)}</b> {x_(b)}", body)) for a, b in R.KNOWN_LIMITS],
                              bulletType="bullet"))
    for cat in CAT_ORDER:
        group = [x for x in rows if x["status"] == cat]
        if not group:
            continue
        story += [PageBreak(), Paragraph(x_(CAT_LABEL[cat]), h2)]
        for x in group:
            story.append(Paragraph(x_(f"{x['id']} · {x['title']}"), h3))
            story.append(Paragraph(x_(x["area"] + (" · You will need: " + x["needs"] if x.get("needs") else "")), small))
            story.append(Paragraph(x_(x["summary"]), body))
            story.append(ListFlowable([ListItem(Paragraph(x_(s), body)) for s in x["howToTest"]], bulletType="1"))
            story.append(Paragraph("<b>We most want to know:</b> " + x_(x["mostWanted"]), body))
            for k, label in (("deliberate", "Looks wrong, is not"), ("knownLimit", "Known limit"),
                             ("weCouldNotVerify", "What our tests cover")):
                if x.get(k):
                    story.append(Paragraph(f"<i>{label}:</i> " + x_(x[k]), small))
            story.append(Spacer(1, 8))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                      bottomMargin=14 * mm, title=R.ROUND["title"]).build(story)


# ---- verification -----------------------------------------------------------------------------------------------

CT_WORKSHEET = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr",
                "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
                "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations",
                "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
                "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
                "legacyDrawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts", "extLst"]


def check_schema_order(xlsx):
    with zipfile.ZipFile(xlsx) as z:
        for name in sorted(n for n in z.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n)):
            root = ET.fromstring(z.read(name))
            last = -1
            for child in root:
                tag = child.tag.split("}")[-1]
                if tag not in CT_WORKSHEET:
                    raise BuildError(f"{name}: unknown element <{tag}>")
                i = CT_WORKSHEET.index(tag)
                if i < last:
                    raise BuildError(f"{name}: <{tag}> out of CT_Worksheet order (Excel would 'repair' and drop it)")
                last = i


def check_workbook(xlsx, rows):
    wb = load_workbook(xlsx)
    want = ["Start here", "What changed", "Known limits", "Test plan", "Test steps", "Anything new", "Progress"]
    if wb.sheetnames != want:
        raise BuildError(f"sheets {wb.sheetnames} != {want}")
    for ws in wb.worksheets:
        for rng in ws.merged_cells.ranges:
            c = ws.cell(row=rng.min_row, column=rng.min_col)
            if c.alignment and c.alignment.wrap_text and rng.min_row == rng.max_row:
                h = ws.row_dimensions[rng.min_row].height
                if not h:
                    raise BuildError(f"{ws.title} row {rng.min_row}: merged wrapping row has no computed height")
        for dv in ws.data_validations.dataValidation:
            f = dv.formula1 or ""
            if len(f) > 255:
                raise BuildError(f"{ws.title}: dropdown formula over 255 characters")
    tp = wb["Test plan"]
    if tp.freeze_panes != "B2" or not tp.auto_filter.ref:
        raise BuildError("Test plan: header row and id column must be frozen, with an autofilter")
    if not tp.conditional_formatting:
        raise BuildError("Test plan: no conditional formatting on Result")
    if not any(isinstance(c.value, str) and c.value.startswith("=COUNTIF") for row in wb["Progress"].iter_rows()
               for c in row):
        raise BuildError("Progress: counts are not formulas")
    if not wb.calculation.fullCalcOnLoad:
        raise BuildError("workbook: fullCalcOnLoad is off")
    text = norm(" ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value))
    return text


def check_coverage(rows, texts):
    for x in rows:
        probes = [x["id"], x["title"]] + [s.split("->")[0] for s in x["howToTest"]] + \
                 [s.split("->")[1] for s in x["howToTest"] if "->" in s]
        for name, t in texts.items():
            for p in probes:
                if norm(p) not in t:
                    raise BuildError(f"{name}: row {x['id']} is missing {p[:60]!r}")


def pdf_text(path):
    import pymupdf
    doc = pymupdf.open(path)
    return norm(" ".join(page.get_text() for page in doc))


def build(rows_in, out_dir):
    rows = checked_rows(copy.deepcopy(rows_in))
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.join(out_dir, f"{R.PRODUCT.lower()}-{R.ROUND['slug']}")
    build_workbook(rows, base + ".xlsx")
    build_csv(rows, base + ".csv")
    build_html(rows, base + ".html")
    build_pdf(rows, base + ".pdf")
    check_schema_order(base + ".xlsx")
    texts = {"xlsx": check_workbook(base + ".xlsx", rows), "pdf": pdf_text(base + ".pdf"),
             "html": norm(html.unescape(re.sub(r"<[^>]+>", " ", open(base + ".html", encoding="utf-8").read()))),
             "csv": norm(open(base + ".csv", encoding="utf-8").read())}
    check_coverage(rows, texts)
    return base, rows


def prove():
    """Break each invariant and confirm the build refuses it (§8: a check that never went red is not a check)."""
    import tempfile
    tmp = tempfile.mkdtemp(prefix="nyar-tp-prove-")
    cases = []

    def expect(name, fn):
        try:
            fn()
        except BuildError as ex:
            cases.append((name, True, str(ex)))
            return
        cases.append((name, False, "no error"))

    bad = copy.deepcopy(R.ROWS); bad[0]["note"] = "x"
    expect("unrendered key", lambda: build(bad, tmp))
    bad = copy.deepcopy(R.ROWS); bad[0]["status"] = "maybe"
    expect("unknown category", lambda: build(bad, tmp))
    bad = copy.deepcopy(R.ROWS); del bad[0]["mostWanted"]
    expect("missing required field", lambda: build(bad, tmp))
    saved = R.TESTERS[:]
    R.TESTERS[0] = "Smith, A"
    expect("comma in a dropdown option", lambda: build(R.ROWS, tmp))
    R.TESTERS[:] = saved

    base, rows = build(R.ROWS, tmp)
    # merged row without height
    wb = load_workbook(base + ".xlsx")
    ws = wb["Start here"]
    ws.row_dimensions[1].height = None
    wb.save(base + "-noheight.xlsx")
    expect("merged prose row without a height", lambda: check_workbook(base + "-noheight.xlsx", rows))
    # schema order: move autoFilter to the end of the Test plan part
    src, dst = base + ".xlsx", base + "-order.xlsx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/worksheets/sheet4.xml":
                s = data.decode("utf-8")
                m = re.search(r"<autoFilter[^>]*/>|<autoFilter.*?</autoFilter>", s, re.S)
                s = s.replace(m.group(0), "").replace("</worksheet>", m.group(0) + "</worksheet>")
                data = s.encode("utf-8")
            zo.writestr(item, data)
    expect("autoFilter out of schema order", lambda: check_schema_order(dst))
    # a step missing from one artefact
    texts = {"pdf": pdf_text(base + ".pdf").replace(norm(rows[0]["howToTest"][1].split("->")[0]), "")}
    expect("a step missing from the PDF", lambda: check_coverage(rows, texts))
    ok = all(c[1] for c in cases)
    for name, went_red, msg in cases:
        print(f"  {'RED (good)' if went_red else 'STAYED GREEN (bad)'}: {name} — {msg[:110]}")
    print(f"prove: {sum(c[1] for c in cases)}/{len(cases)} invariants went red")
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)
    return 0 if ok else 1


if __name__ == "__main__":
    if "--prove" in sys.argv:
        sys.exit(prove())
    base, rows = build(R.ROWS, OUT)
    steps = sum(len(x["howToTest"]) for x in rows)
    print(f"built {len(rows)} rows, {steps} steps -> {base}.xlsx / .pdf / .csv / .html; verified (coverage in all four, "
          f"schema order, merged heights, dropdowns, formulas)")
